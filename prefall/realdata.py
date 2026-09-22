"""Loaders that turn real wearable datasets into the same `Trial` objects the simulator produces.

Supported (formats taken from each dataset's own documentation):

  SisFall  waist, 200 Hz, files  <CODE>_<SUBJECT>_<TRIAL>.txt  (e.g. F05_SA01_R04.txt), 9 columns of raw
           ADC counts: ADXL345 accel, ITG3200 gyro, MMA8451Q accel. No fall-timing labels.
  KFall    low back, CSV per motion file  <SUBJECT>T<TASK>R<TRIAL>.csv  (e.g. SA06T20R01.csv) with
           time, frame counter, accel (g), gyro (deg/s), Euler angles; Excel label files with the fall
           onset and impact frame of every fall trial.

Every trial comes out in the algorithm's frame (see synth.py): channels ax ay az [g], gx gy gz [deg/s],
az = +1 g when upright, and gy = +d(angle)/dt for ax = +sin(angle). Neither dataset documents which physical
axis points up, so `estimate_axis_map` works it out from the recordings themselves and reports how confident it is.
"""
from __future__ import annotations

import csv
import glob
import os
import re
import warnings
from dataclasses import dataclass, field

import numpy as np
from scipy import signal

from .synth import Trial

# ------------------------------------------------------------------------------------------------
# axis mapping
# ------------------------------------------------------------------------------------------------


@dataclass
class AxisMap:
    """For each output channel: which native axis feeds it, and with which sign."""
    acc_src: tuple = (0, 1, 2)
    acc_sign: tuple = (1.0, 1.0, 1.0)
    gyr_src: tuple = (0, 1, 2)
    gyr_sign: tuple = (1.0, 1.0, 1.0)
    info: dict = field(default_factory=dict)

    def apply(self, acc, gyr):
        a = acc[:, list(self.acc_src)] * np.array(self.acc_sign)
        g = gyr[:, list(self.gyr_src)] * np.array(self.gyr_sign)
        return np.hstack([a, g])

    def describe(self):
        names = ["x", "y", "z"]
        a = ", ".join(f"{o}<-{'-' if s < 0 else '+'}{names[i]}" for o, i, s in zip(["ax", "ay", "az"], self.acc_src, self.acc_sign))
        g = ", ".join(f"{o}<-{'-' if s < 0 else '+'}{names[i]}" for o, i, s in zip(["gx", "gy", "gz"], self.gyr_src, self.gyr_sign))
        return f"accel: {a} | gyro: {g}"


def _lowpass(x, fs, fc=5.0):
    b, a = signal.butter(2, min(fc / (fs / 2.0), 0.99))
    if len(x) < 30:
        return x
    return signal.filtfilt(b, a, x, axis=0)


def estimate_axis_map(samples, rest_s=1.0):
    """samples: list of (acc [n,3] in g, gyr [n,3] in deg/s, fs), in the dataset's native axes.

    1. Vertical axis: the axis with the largest median reading over the first second of many recordings
       (people start upright); its sign makes az = +1 g.
    2. Forward axis (ax): of the two horizontal axes, the one that swings more across the whole dataset.
       This is a heuristic (it assumes forward/backward bending, sitting and falls dominate over sideways
       tilt), and it is printed so you can check it.
    3. Gyro axes: for each horizontal accelerometer axis, the gyro channel (and sign) whose signal best
       matches the time derivative of the tilt angle computed from the accelerometer.
    """
    rest = np.array([acc[:max(2, int(rest_s * fs))].mean(0) for acc, _, fs in samples])
    m = np.median(rest, axis=0)
    v = int(np.argmax(np.abs(m)))
    sv = 1.0 if m[v] > 0 else -1.0
    if not 0.7 < abs(m[v]) < 1.3:
        warnings.warn(f"vertical axis reads {abs(m[v]):.2f} g at rest (expected about 1 g): check the units")

    others = [i for i in range(3) if i != v]
    spread = np.zeros(3)
    for acc, _, fs in samples:
        spread += _lowpass(acc, fs).var(axis=0)
    h1 = others[int(np.argmax(spread[others]))]
    h2 = [i for i in others if i != h1][0]
    acc_src, acc_sign = (h1, h2, v), (1.0, 1.0, sv)

    # gyro matching
    sxy = np.zeros((2, 3))
    sxx, syy = np.zeros(2), np.zeros(3)
    for acc, gyr, fs in samples:
        a3 = acc[:, list(acc_src)] * np.array(acc_sign)
        anorm = np.linalg.norm(a3, axis=1)
        ok = (anorm > 0.85) & (anorm < 1.15)
        if ok.sum() < 20:
            continue
        af = _lowpass(a3, fs)
        for k, h in enumerate((0, 1)):               # k = 0: sagittal tilt from ax, k = 1: lateral tilt from ay
            th = np.degrees(np.arctan2(af[:, h], af[:, 2]))
            d = np.gradient(th) * fs
            gf = _lowpass(gyr, fs)
            sxx[k] += (d[ok] ** 2).sum()
            sxy[k] += (d[ok, None] * gf[ok]).sum(0)
        syy += (_lowpass(gyr, fs)[ok] ** 2).sum(0)
    corr = sxy / np.sqrt(np.maximum(sxx[:, None] * syy[None, :], 1e-12))
    j_gy = int(np.argmax(np.abs(corr[0])))
    s_gy = 1.0 if corr[0, j_gy] > 0 else -1.0
    rest_j = [j for j in range(3) if j != j_gy]
    j_gx = rest_j[int(np.argmax(np.abs(corr[1, rest_j])))]
    s_gx = 1.0 if corr[1, j_gx] > 0 else -1.0
    j_gz = [j for j in range(3) if j not in (j_gy, j_gx)][0]
    amap = AxisMap(acc_src, acc_sign, (j_gx, j_gy, j_gz), (s_gx, s_gy, 1.0))
    amap.info = dict(rest_vector=m.round(3).tolist(), horizontal_spread=spread.round(3).tolist(),
                     corr_sagittal=corr[0].round(2).tolist(), corr_lateral=corr[1].round(2).tolist(),
                     n_recordings=len(samples))
    if abs(corr[0, j_gy]) < 0.3:
        warnings.warn(f"gyro/accelerometer match is weak (correlation {abs(corr[0, j_gy]):.2f}): "
                      "the gyro may be in a different frame; the threshold detector will be unreliable")
    return amap


# ------------------------------------------------------------------------------------------------
# SisFall
# ------------------------------------------------------------------------------------------------

_SIS_RE = re.compile(r"^([DF]\d{2})_(S[AE]\d{2})_R(\d{2})\.txt$", re.I)
SIS_ACC_G_PER_LSB = 2 * 16 / 2 ** 13        # ADXL345: +-16 g, 13 bit (dataset readme formula)
SIS_GYR_DPS_PER_LSB = 2 * 2000 / 2 ** 16    # ITG3200: +-2000 deg/s, 16 bit
SIS_FS = 200.0


def parse_sisfall_file(path):
    """Returns an (n, 9) array of raw counts. Tolerates the trailing ';' on each line and stray whitespace."""
    with open(path, "r", errors="ignore") as f:
        text = f.read().replace(";", " ").replace(",", " ")
    vals = np.array(text.split(), dtype=float)
    if len(vals) % 9:
        warnings.warn(f"{os.path.basename(path)}: {len(vals)} numbers is not a multiple of 9, truncating")
        vals = vals[:len(vals) // 9 * 9]
    return vals.reshape(-1, 9)


def load_sisfall(root, limit_subjects=None, onset_window=1.0, min_impact_g=2.0, map_sample=80, seed=0, verbose=True):
    """Load SisFall from a folder that contains the subject folders (SA01 ... SE15) with the .txt files.

    Fall timing: SisFall has no labels, so the impact is ESTIMATED as the sample with the largest acceleration
    magnitude, and the fall onset is assumed to be `onset_window` seconds earlier. Falls whose peak is below
    `min_impact_g` are dropped (no clear impact to anchor on). Treat lead times as approximate.
    """
    files = sorted(p for p in glob.glob(os.path.join(root, "**", "*.txt"), recursive=True)
                   if _SIS_RE.match(os.path.basename(p)))
    if not files:
        raise FileNotFoundError(f"no SisFall files (like F05_SA01_R04.txt) found under {root}")
    subjects = sorted({_SIS_RE.match(os.path.basename(p)).group(2).upper() for p in files})
    if limit_subjects:
        keep = set(subjects[:limit_subjects])
        files = [p for p in files if _SIS_RE.match(os.path.basename(p)).group(2).upper() in keep]
    raws = []
    for k, p in enumerate(files):
        code, subj, trial = _SIS_RE.match(os.path.basename(p)).groups()
        a9 = parse_sisfall_file(p)
        raws.append((code.upper(), subj.upper(), int(trial), a9[:, :3] * SIS_ACC_G_PER_LSB, a9[:, 3:6] * SIS_GYR_DPS_PER_LSB))
        if verbose and (k + 1) % 500 == 0:
            print(f"  read {k + 1}/{len(files)} SisFall files")
    rng = np.random.default_rng(seed)
    pick = rng.choice(len(raws), size=min(map_sample, len(raws)), replace=False)
    amap = estimate_axis_map([(raws[i][3], raws[i][4], SIS_FS) for i in pick])
    trials, dropped = [], 0
    for code, subj, trial, acc, gyr in raws:
        x = amap.apply(acc, gyr)
        t_on = t_imp = None
        if code.startswith("F"):
            n = np.linalg.norm(x[:, :3], axis=1)
            i = int(np.argmax(n))
            if n[i] < min_impact_g:
                dropped += 1
                continue
            t_imp = i / SIS_FS
            t_on = max(0.0, t_imp - onset_window)
        trials.append(Trial("fall" if code.startswith("F") else code, x, t_on, t_imp,
                            dict(dataset="sisfall", id=subj, task=code, trial=trial,
                                 group="elderly" if subj.startswith("SE") else "adult"), SIS_FS))
    if verbose:
        nf = sum(t.kind == "fall" for t in trials)
        print(f"SisFall: {len(trials)} recordings ({nf} falls, {len(trials) - nf} activities of daily living), "
              f"{len({t.subject['id'] for t in trials})} subjects; {dropped} falls dropped (no clear impact)")
        print(f"  axis map: {amap.describe()}")
        print(f"  resting gravity vector (native axes, g): {amap.info['rest_vector']}, "
              f"gyro match (sagittal): {amap.info['corr_sagittal']}")
    return trials, amap


# ------------------------------------------------------------------------------------------------
# KFall
# ------------------------------------------------------------------------------------------------

_KF_RE = re.compile(r"^(S[A-Z]?\d{2})T(\d{2})R(\d{2})\.csv$", re.I)

def _norm_kfall_subj(s):
    """Some KFall drops name sensor files 'S06T...' but label files/folders 'SA06' (or vice versa).
    Normalize both to 'SA<NN>' so a fall recording actually finds its label."""
    m = re.search(r"(\d{2,3})", s)
    return f"SA{int(m.group(1)):02d}" if m else s.upper()

def parse_kfall_csv(path):
    """Returns (time_s, frame, acc[n,3] g, gyr[n,3] deg/s). Columns are read by position (time, frame counter,
    3 accel, 3 gyro, 3 Euler) and non-numeric header rows are skipped, so header wording does not matter."""
    rows = []
    with open(path, "r", newline="", encoding="utf-8-sig", errors="ignore") as f:
        for r in csv.reader(f):
            if len(r) < 8:
                continue
            try:
                rows.append([float(c) for c in r[:8]])
            except ValueError:
                continue                        # header
    a = np.array(rows)
    if len(a) == 0:
        raise ValueError(f"{path}: no numeric rows")
    return a[:, 0], a[:, 1], a[:, 2:5], a[:, 5:8]


def _task_id(cell):
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return int(cell)
    s = str(cell)
    m = re.search(r"\((\d+)\)", s) or re.fullmatch(r"\s*(\d+)\s*", s)
    return int(m.group(1)) if m else None


def read_kfall_labels(label_dir):
    """{(subject, task_id, trial_id): (onset_frame, impact_frame)} from the *_label.xlsx files."""
    try:
        import openpyxl
    except ImportError as e:                    # pragma: no cover
        raise ImportError("reading KFall labels needs openpyxl: pip install openpyxl") from e
    labels, bad = {}, 0
    files = sorted(glob.glob(os.path.join(label_dir, "**", "*.xlsx"), recursive=True))
    if not files:
        raise FileNotFoundError(f"no .xlsx label files found under {label_dir}")
    for f in files:
        m = re.search(r"(S[A-Z]?\d{2})", os.path.basename(f), re.I)
        if not m:
            continue
        subj = _norm_kfall_subj(m.group(1))        
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        head = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and "onset" in c.lower() for c in r)), None)
        if head is None:
            warnings.warn(f"{os.path.basename(f)}: no header row containing 'onset', skipped")
            continue
        cols = {}
        for j, c in enumerate(rows[head]):
            c = str(c).lower() if c is not None else ""
            if "task" in c:
                cols["task"] = j
            elif "trial" in c:
                cols["trial"] = j
            elif "onset" in c:
                cols["onset"] = j
            elif "impact" in c:
                cols["impact"] = j
        if len(cols) < 4:
            warnings.warn(f"{os.path.basename(f)}: expected task/trial/onset/impact columns, found {sorted(cols)}")
            continue
        task = None
        for r in rows[head + 1:]:
            t = _task_id(r[cols["task"]]) if r[cols["task"]] is not None else None
            task = t if t is not None else task     # merged cells: the task id is only on the first row of a group
            try:
                trial, on, imp = int(r[cols["trial"]]), float(r[cols["onset"]]), float(r[cols["impact"]])
            except (TypeError, ValueError):
                continue
            if task is None:
                bad += 1
                continue
            labels[(subj, task, trial)] = (on, imp)
    if bad:
        warnings.warn(f"{bad} label rows had no readable task id")
    return labels


def load_kfall(root, limit_subjects=None, map_sample=80, seed=0, verbose=True):
    """Load KFall. `root` holds sensor_data/ and label_data/ (the two unzipped downloads)."""
    sensor_dir, label_dir = os.path.join(root, "sensor_data"), os.path.join(root, "label_data")
    for d in (sensor_dir, label_dir):
        if not os.path.isdir(d):
            raise FileNotFoundError(f"{d} not found: unzip sensor_data.zip and label_data.zip into {root}")
    labels = read_kfall_labels(label_dir)
    fall_tasks = {t for (_, t, _) in labels}
    files = sorted(p for p in glob.glob(os.path.join(sensor_dir, "**", "*.csv"), recursive=True)
                   if _KF_RE.match(os.path.basename(p)))
    if not files:
        raise FileNotFoundError(f"no KFall files (like SA06T20R01.csv) found under {sensor_dir}")
    if limit_subjects:
        subs = sorted({_norm_kfall_subj(_KF_RE.match(os.path.basename(p)).group(1)) for p in files})[:limit_subjects]
        files = [p for p in files if _norm_kfall_subj(_KF_RE.match(os.path.basename(p)).group(1)) in subs]
    raws = []
    for k, p in enumerate(files):
        subj, task, trial = _KF_RE.match(os.path.basename(p)).groups()
        t, frame, acc, gyr = parse_kfall_csv(p)
        dt = float(np.median(np.diff(t))) if len(t) > 2 else 0.0
        fs = 1.0 / dt if dt > 0 and 20 < 1.0 / dt < 1000 else 100.0
        raws.append((_norm_kfall_subj(subj), int(task), int(trial), frame, acc, gyr, round(fs, 2)))
        if verbose and (k + 1) % 500 == 0:
            print(f"  read {k + 1}/{len(files)} KFall files")
    rng = np.random.default_rng(seed)
    pick = rng.choice(len(raws), size=min(map_sample, len(raws)), replace=False)
    amap = estimate_axis_map([(raws[i][4], raws[i][5], raws[i][6]) for i in pick])
    trials, unlabeled = [], 0
    for subj, task, trial, frame, acc, gyr, fs in raws:
        x = amap.apply(acc, gyr)
        lab = labels.get((subj, task, trial))
        if lab is not None:
            i_on = int(np.searchsorted(frame, lab[0]))
            i_im = int(np.searchsorted(frame, lab[1]))
            if not 0 <= i_on < i_im <= len(x):
                warnings.warn(f"{subj}T{task:02d}R{trial:02d}: label frames outside the recording, skipped")
                continue
            trials.append(Trial("fall", x, i_on / fs, i_im / fs,
                                dict(dataset="kfall", id=subj, task=f"T{task:02d}", trial=trial, group="adult"), fs))
        elif task in fall_tasks:
            unlabeled += 1                      # a fall task without its own label: cannot be used as an ADL
        else:
            trials.append(Trial(f"T{task:02d}", x, None, None,
                                dict(dataset="kfall", id=subj, task=f"T{task:02d}", trial=trial, group="adult"), fs))
    if verbose:
        nf = sum(t.kind == "fall" for t in trials)
        fss = sorted({t.fs_raw for t in trials})
        print(f"KFall: {len(trials)} recordings ({nf} falls, {len(trials) - nf} activities of daily living), "
              f"{len({t.subject['id'] for t in trials})} subjects, sampling rate(s) {fss} Hz; "
              f"{unlabeled} fall recordings without a label skipped")
        print(f"  axis map: {amap.describe()}")
        print(f"  resting gravity vector (native axes, g): {amap.info['rest_vector']}, "
              f"gyro match (sagittal): {amap.info['corr_sagittal']}")
    return trials, amap


# ------------------------------------------------------------------------------------------------
# helpers
# ------------------------------------------------------------------------------------------------

def split_subjects(trials, test_frac=0.3, seed=0, min_test_falls=5):
    """Subject-wise split: no person appears in both train and test (a must for wearable data)."""
    ids = sorted({t.subject["id"] for t in trials})
    rng = np.random.default_rng(seed)
    for _ in range(200):
        order = list(rng.permutation(ids))
        n_test = max(1, int(round(test_frac * len(ids))))
        test_ids = set(order[:n_test])
        test = [t for t in trials if t.subject["id"] in test_ids]
        if sum(t.kind == "fall" for t in test) >= min_test_falls:
            break
    train = [t for t in trials if t.subject["id"] not in test_ids]
    return train, test


def inspect(root, dataset):
    """Print the first lines of a few raw files, so a format surprise can be diagnosed by eye."""
    if dataset == "sisfall":
        files = sorted(p for p in glob.glob(os.path.join(root, "**", "*.txt"), recursive=True) if _SIS_RE.match(os.path.basename(p)))[:2]
        for p in files:
            print(f"--- {p}")
            with open(p, errors="ignore") as f:
                for _ in range(3):
                    print(repr(f.readline()))
    else:
        files = sorted(glob.glob(os.path.join(root, "sensor_data", "**", "*.csv"), recursive=True))[:1]
        for p in files:
            print(f"--- {p}")
            with open(p, errors="ignore") as f:
                for _ in range(4):
                    print(repr(f.readline()))
        try:
            import openpyxl
            for p in sorted(glob.glob(os.path.join(root, "label_data", "**", "*.xlsx"), recursive=True))[:1]:
                print(f"--- {p}")
                for r in list(openpyxl.load_workbook(p, read_only=True, data_only=True).active.iter_rows(values_only=True))[:6]:
                    print(r)
        except ImportError:                     # pragma: no cover
            print("(install openpyxl to inspect the label files)")
